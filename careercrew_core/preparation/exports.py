"""Render complete persisted resume text; no HTML, URL fetching, or temp files."""
from io import BytesIO


def export_pdf(label: str, content: str) -> bytes:
    import fitz

    font = fitz.Font("cjk")  # bundled CJK font; independent of host font installs
    size, line_height, margin = 11, 18, 48
    width, height = 595, 842
    advance = {}
    with fitz.open() as document:
        page = None
        y = height

        def put_line(text):
            nonlocal page, y
            if y + line_height > height - margin:
                page = document.new_page(width=width, height=height)
                page.insert_font(fontname="resume", fontbuffer=font.buffer)
                y = margin + size
            if text:
                page.insert_text((margin, y), text, fontname="resume", fontsize=size)
            y += line_height

        # Wrap by glyph width, so both unbroken Chinese and long URLs fit.
        for paragraph in (label + "\n\n" + content).split("\n"):
            line, used = [], 0.0
            for char in paragraph.rstrip("\r").expandtabs(4):
                if char not in advance:
                    advance[char] = font.glyph_advance(ord(char)) * size
                amount = advance[char]
                if line and used + amount > width - 2 * margin:
                    put_line("".join(line))
                    line, used = [], 0.0
                line.append(char)
                used += amount
            put_line("".join(line))
        document.subset_fonts()
        return document.tobytes(garbage=4, deflate=True)


def export_docx(label: str, content: str) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.element.get_or_add_rPr().append(fonts)
    document.add_heading(label, level=1)
    for paragraph in content.split("\n"):
        document.add_paragraph(paragraph.rstrip("\r"))
    output = BytesIO()
    document.save(output)
    return output.getvalue()
